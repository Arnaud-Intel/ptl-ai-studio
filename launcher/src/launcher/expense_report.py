"""Durable review state for extracted receipts, independent of the model workers."""
from __future__ import annotations

import io
import json
import re
import sqlite3
from contextlib import contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from expense_extract.parsing import CURRENCIES, coerce_amount

from .errors import Conflict

CATEGORIES = ("Meals", "Travel", "Lodging", "Office Supplies", "Software", "Other")
FIELDS = ("vendor", "date", "amount", "currency", "category", "notes")


class ExpenseReports:
    def __init__(self, path: Path | None = None):
        self.path = path or Path(__file__).resolve().parents[3] / "logs" / "expense-reports.sqlite3"

    @contextmanager
    def _connect(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        db = sqlite3.connect(self.path, timeout=15)
        db.execute("CREATE TABLE IF NOT EXISTS reports (id TEXT PRIMARY KEY, created TEXT, folder TEXT)")
        db.execute("CREATE TABLE IF NOT EXISTS expenses (id TEXT PRIMARY KEY, report TEXT, data TEXT)")
        try:
            with db:
                yield db
        finally:
            db.close()

    def create(self, folder: str) -> str:
        report_id = uuid4().hex
        with self._connect() as db:
            db.execute("INSERT INTO reports VALUES (?, ?, ?)",
                       (report_id, datetime.now(timezone.utc).isoformat(), str(Path(folder).resolve())))
        return report_id

    def add(self, report_id: str, line) -> dict:
        original = line.to_dict()
        item = {**original, "id": uuid4().hex, "report_id": report_id, "revision": 0,
                "status": "draft", "notes": "", "validated_at": None, "original": original}
        with self._connect() as db:
            db.execute("INSERT INTO expenses VALUES (?, ?, ?)", (item["id"], report_id, json.dumps(item)))
        return item

    def snapshot(self, report_id: str | None = None) -> dict:
        with self._connect() as db:
            reports = [{"id": r[0], "created": r[1], "folder": r[2]} for r in
                       db.execute("SELECT id, created, folder FROM reports ORDER BY created DESC")]
            selected = next((r for r in reports if r["id"] == report_id), None) if report_id else next(iter(reports), None)
            if report_id and selected is None:
                raise ValueError("Report not found.")
            items = [json.loads(r[0]) for r in db.execute(
                "SELECT data FROM expenses WHERE report = ? ORDER BY rowid", (selected["id"],))] if selected else []
        totals = {}
        for item in items:
            if item["status"] == "validated":
                currency = item["currency"]
                totals[currency] = totals.get(currency, Decimal(0)) + Decimal(item["amount"])
        return {"report": selected, "reports": reports, "items": items,
                "totals": {k: format(v, ".2f") for k, v in sorted(totals.items())},
                "currencies": sorted(CURRENCIES), "categories": CATEGORIES}

    def update(self, report_id: str, item_id: str, values: dict, revision: int, validate: bool) -> dict:
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            row = db.execute("SELECT data FROM expenses WHERE id = ? AND report = ?", (item_id, report_id)).fetchone()
            if row is None:
                raise ValueError("Expense not found.")
            item = json.loads(row[0])
            if item["revision"] != revision:
                raise Conflict("This expense changed in another view. Reopen it before saving.")
            item.update({key: str(values.get(key) or "").strip() for key in FIELDS})
            if any(len(item[key]) > 4000 for key in FIELDS):
                raise ValueError("Expense fields must be at most 4,000 characters.")
            item["currency"] = item["currency"].upper()
            amount = coerce_amount(item["amount"])
            reasons = []
            if not item["vendor"]:
                reasons.append("Enter a vendor")
            try:
                if date.fromisoformat(item["date"]).isoformat() != item["date"]:
                    raise ValueError
            except ValueError:
                reasons.append("Enter a valid date (YYYY-MM-DD)")
            if amount is None:
                reasons.append("Enter an unambiguous amount with at most two decimals")
            if item["currency"] not in CURRENCIES:
                reasons.append("Choose a supported currency")
            if amount is not None and item["currency"] in {"JPY", "KRW"} and amount != amount.to_integral_value():
                reasons.append("This currency requires a whole-number amount")
            if item["category"] not in CATEGORIES:
                reasons.append("Choose a category")
            if validate and reasons:
                raise ValueError("; ".join(reasons) + ".")
            if amount is not None:
                item["amount"] = format(amount, ".2f")
            item.update(status="validated" if validate else "draft", needs_review=not validate,
                        review_reasons=reasons, error=None, revision=revision + 1,
                        validated_at=datetime.now(timezone.utc).isoformat() if validate else None)
            db.execute("UPDATE expenses SET data = ? WHERE id = ?", (json.dumps(item), item_id))
        return item

    def receipt(self, report_id: str, item_id: str) -> Path:
        report = self.snapshot(report_id)
        item = next((i for i in report["items"] if i["id"] == item_id), None)
        if item is None:
            raise ValueError("Expense not found.")
        folder = Path(report["report"]["folder"]).resolve()
        path = (folder / item["source_file"]).resolve()
        if path.parent != folder or not path.is_file():
            raise FileNotFoundError("The original receipt is no longer available in its source folder.")
        return path

    def export(self, report_id: str) -> bytes:
        # A normal installable Python dependency: export also works outside Codex.
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        report = self.snapshot(report_id)
        if not report["items"]:
            raise Conflict("Read at least one receipt before exporting.")
        wb = Workbook()
        summary = wb.active
        summary.title = "Summary"
        summary.append(["Expense report", report["report"]["created"][:10]])
        summary.append(["All expenses", len(report["items"])])
        summary.append(["Validated", sum(i["status"] == "validated" for i in report["items"])])
        summary.append(["Drafts excluded from totals", sum(i["status"] != "validated" for i in report["items"])])
        summary.append([])
        summary.append(["Currency", "Validated total"])
        details = wb.create_sheet("Expenses")
        details.append(["Receipt", "Vendor", "Date", "Amount", "Currency", "Category", "Status", "Notes", "Validated at"])
        source = wb.create_sheet("Extraction")
        source.append(["Receipt", "Original vendor", "Original date", "Original amount", "Original currency",
                       "Original category", "Extraction warnings", "Receipt text"])
        def append_receipt_row(sheet, values):
            # OCR may contain control characters that XML cannot represent.
            sheet.append([re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
                          if isinstance(value, str) else value for value in values])

        for item in report["items"]:
            amount = coerce_amount(item["amount"])
            try:
                expense_date = date.fromisoformat(item["date"])
            except ValueError:
                expense_date = item["date"]
            append_receipt_row(details, [item["source_file"], item["vendor"], expense_date,
                            float(amount) if amount is not None else item["amount"], item["currency"],
                            item["category"], item["status"].title(), item["notes"], item["validated_at"] or ""])
            original = item["original"]
            append_receipt_row(source, [original.get(k) for k in ("source_file", "vendor", "date", "amount", "currency", "category")] +
                          ["; ".join(filter(None, [original.get("error"), *original.get("review_reasons", [])])), original["raw_text"]])
        # Force all supplied strings to text, including strings starting with '='.
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="17324D")
                cell.font = Font(color="FFFFFF", bold=True)
            for column in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(column)].width = 22
        end = len(report["items"]) + 1
        for row, currency in enumerate(report["totals"], start=7):
            summary.cell(row, 1, currency)
            summary.cell(row, 2, f'=SUMIFS(Expenses!D2:D{end},Expenses!E2:E{end},A{row},Expenses!G2:G{end},"Validated")').number_format = '#,##0.00;[Red](#,##0.00)'
        summary.column_dimensions["A"].width = 34
        for sheet in (details, source):
            sheet.auto_filter.ref = sheet.dimensions
            sheet.column_dimensions["A"].width = 30
            sheet.column_dimensions["H"].width = 60
        for row in range(2, end + 1):
            details.cell(row, 3).number_format = "yyyy-mm-dd"
            details.cell(row, 4).number_format = '#,##0.00;[Red](#,##0.00)'
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

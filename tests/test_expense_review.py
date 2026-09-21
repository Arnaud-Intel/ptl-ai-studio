from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from expense_extract.types import ExpenseLine
from launcher.errors import Conflict
from launcher.expense_report import ExpenseReports


@pytest.fixture
def report(tmp_path):
    store = ExpenseReports(tmp_path / "reports.sqlite3")
    report_id = store.create(str(tmp_path))
    item = store.add(report_id, ExpenseLine("receipt.png", "Cafe", "2026-09-14", Decimal("12.50"),
                                         "Meals", "Cafe total EUR 12.50", currency="EUR"))
    return store, report_id, item


def fields(**changes):
    return dict(vendor="Cafe", date="2026-09-14", amount="12.50", currency="EUR", category="Meals", notes="Lunch", **changes)


def test_review_is_explicit_durable_and_edits_reopen_validation(report):
    store, rid, item = report
    assert store.snapshot(rid)["totals"] == {}
    saved = store.update(rid, item["id"], fields(), 0, True)
    assert saved["validated_at"]
    restored = ExpenseReports(store.path).snapshot(rid)
    assert restored["totals"] == {"EUR": "12.50"}
    values = fields()
    values["amount"] = "13,75"
    draft = store.update(rid, item["id"], values, 1, False)
    assert draft["amount"] == "13.75"
    assert draft["original"]["amount"] == "12.50"
    assert draft["validated_at"] is None
    assert store.snapshot(rid)["totals"] == {}
    with pytest.raises(Conflict):
        store.update(rid, item["id"], values, 1, True)


@pytest.mark.parametrize("key,value", [("amount", "1,234"), ("amount", "NaN"), ("date", "2026-02-30"),
                                     ("vendor", ""), ("currency", "???"), ("category", "Invalid")])
def test_invalid_values_can_be_drafted_but_not_validated(report, key, value):
    store, rid, item = report
    values = fields()
    values[key] = value
    with pytest.raises(ValueError):
        store.update(rid, item["id"], values, 0, True)
    assert store.update(rid, item["id"], values, 0, False)["status"] == "draft"


def test_failed_extraction_can_be_completed_and_reports_stay_separate(report):
    store, rid, _ = report
    other = store.create(".")
    item = store.add(other, ExpenseLine("bad.png", "", "", None, "Other", "", error="OCR failed"))
    saved = store.update(other, item["id"], fields(), 0, True)
    assert saved["error"] is None
    assert saved["original"]["error"] == "OCR failed"
    assert len(store.snapshot(rid)["items"]) == 1
    assert len(store.snapshot()["reports"]) == 2
    with pytest.raises(ValueError):
        store.update(rid, item["id"], fields(), 1, False)


def test_excel_has_all_rows_real_numbers_dates_and_safe_text(report):
    store, rid, item = report
    values = fields()
    values["vendor"] = '=HYPERLINK("https://example.com", "click")'
    store.update(rid, item["id"], values, 0, True)
    store.add(rid, ExpenseLine("other.png", "Other", "", None, "Other", "=1+1\x00", error="Incomplete"))
    wb = load_workbook(BytesIO(store.export(rid)))
    assert wb.sheetnames == ["Summary", "Expenses", "Extraction"]
    assert wb["Expenses"].max_row == 3
    assert wb["Expenses"]["B2"].data_type == "s"
    assert wb["Expenses"]["D2"].value == 12.5
    assert wb["Expenses"]["C2"].value.date().isoformat() == "2026-09-14"
    assert wb["Expenses"]["G3"].value == "Draft"
    assert wb["Extraction"]["H3"].data_type == "s"
    assert wb["Extraction"]["H3"].value == "=1+1"
    assert wb["Summary"]["B7"].value == '=SUMIFS(Expenses!D2:D3,Expenses!E2:E3,A7,Expenses!G2:G3,"Validated")'
    assert wb["Expenses"].freeze_panes == "A2"


def test_preview_cannot_escape_the_report_folder(report):
    store, rid, _ = report
    item = store.add(rid, ExpenseLine("../outside.png", "", "", None, "Other", ""))
    with pytest.raises(FileNotFoundError):
        store.receipt(rid, item["id"])


def test_review_api_and_export(report, monkeypatch):
    from fastapi.testclient import TestClient
    from launcher import app as launcher_app

    store, rid, item = report
    monkeypatch.setattr(launcher_app.expense_extract_runner, "reports", store)
    client = TestClient(launcher_app.app)
    assert client.get("/api/expense-extract/report").json()["items"][0]["status"] == "draft"
    url = f"/api/expense-extract/reports/{rid}/expenses/{item['id']}"
    assert client.put(url, json={**fields(), "revision": 0, "validate": True}).status_code == 200
    assert client.put(url, json={**fields(), "revision": 0, "validate": True}).status_code == 409
    exported = client.get(f"/api/expense-extract/reports/{rid}/export.xlsx")
    assert exported.status_code == 200
    assert exported.content[:2] == b"PK"
    assert "spreadsheetml" in exported.headers["content-type"]


def test_extraction_runner_saves_each_receipt_before_notifying(tmp_path, monkeypatch):
    import asyncio
    from launcher.expense_extract_runner import ExpenseExtractRunner
    from launcher import expense_extract_runner as module
    from pantherlake_ai_core.engine import Engine

    runner = ExpenseExtractRunner()
    runner.reports = ExpenseReports(tmp_path / "reports.sqlite3")
    monkeypatch.setattr(module.pipeline, "list_receipt_images", lambda folder: [tmp_path / "receipt.png"])
    line = ExpenseLine("receipt.png", "Cafe", "2026-09-14", Decimal("12.50"), "Meals", "EUR 12.50", currency="EUR")

    def fake_run(**kwargs):
        kwargs["on_structured"](line)
        return [line]

    monkeypatch.setattr(module.pipeline, "run", fake_run)

    async def run():
        queue = asyncio.Queue()
        runner.start(loop=asyncio.get_running_loop(), queue=queue, folder=str(tmp_path),
                     ocr_engine=Engine.PORTABLE, ocr_device="cpu", llm_engine=Engine.PORTABLE, llm_device="cpu")
        message = await asyncio.wait_for(queue.get(), 5)
        assert message["type"] == "structured"
        snapshot = runner.reports.snapshot(runner.report_id)
        assert snapshot["items"][0]["id"] == message["line"]["id"]
        assert snapshot["items"][0]["status"] == "draft"
        await asyncio.to_thread(runner._thread.join, 5)

    asyncio.run(run())
